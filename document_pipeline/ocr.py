import io
import logging
import re
from pathlib import Path

logger = logging.getLogger("ocr")

SUPPORTED = {
    "text": {".txt", ".md", ".html", ".htm", ".xml", ".csv", ".json", ".log", ".py", ".js", ".yaml", ".yml"},
    "pdf": {".pdf"},
    "docx": {".docx", ".doc"},
    "image": {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".gif", ".webp"},
}


def extract_text(file_path: str, ext: str) -> dict:
    ext = ext.lower().strip()
    try:
        if ext in SUPPORTED["text"]:
            return _read_text_file(file_path)
        if ext in SUPPORTED["pdf"]:
            return _extract_pdf(file_path)
        if ext in SUPPORTED["docx"]:
            return _extract_docx(file_path)
        if ext in SUPPORTED["image"]:
            return _extract_image(file_path)
    except Exception as exc:
        logger.warning("Extraction failed for %s: %s", file_path, exc)
        return _fail(f"Extraction failed: {exc}")
    return _fail(f"Unsupported file type: {ext}")


def _read_text_file(file_path: str) -> dict:
    content = ""
    for enc in ("utf-8", "latin-1", "cp1252", "ascii"):
        try:
            content = Path(file_path).read_text(encoding=enc)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    cleaned = _strip_html_tags(content) if _looks_html(content) else content
    word_count = len(cleaned.split())
    return {
        "text": cleaned,
        "pages": 1,
        "confidence": 95.0 if word_count > 50 else 70.0,
        "engine": "plaintext",
        "warnings": [] if word_count > 50 else ["Very short document"],
    }


def _looks_html(text: str) -> bool:
    sample = text[:500].strip()
    return bool(re.search(r"<(html|body|div|p|table|head)\b", sample, re.IGNORECASE))


def _strip_html_tags(html: str) -> str:
    text = re.sub(r"<style[^>]*>.*?</style>", "", html, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<script[^>]*>.*?</script>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"&nbsp;", " ", text)
    text = re.sub(r"&[a-z]+;", " ", text)
    return " ".join(text.split())


def _extract_pdf(file_path: str) -> dict:
    result = _pdfplumber_extract(file_path)
    if result["text"].strip():
        return result
    ocr_result = _ocr_pdf_pages(file_path)
    if ocr_result["text"].strip():
        return ocr_result
    return _fail("PDF contained no readable text")


def _pdfplumber_extract(file_path: str) -> dict:
    try:
        import pdfplumber
    except ImportError:
        return _fail("pdfplumber not installed: pip install pdfplumber")

    pages_text = []
    confidences = []
    tables_found = 0
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                txt = (page.extract_text() or "").strip()
                pages_text.append(txt)
                confidences.append(95.0 if txt else 0.0)
                tables = page.extract_tables()
                if tables:
                    tables_found += len(tables)
                    for table in tables:
                        for row in table:
                            cells = [str(c).strip() for c in row if c]
                            if cells:
                                pages_text.append(" | ".join(cells))
    except Exception as exc:
        logger.warning("pdfplumber failed: %s", exc)
        return _fail(f"pdfplumber error: {exc}")

    full = "\n\n".join(p for p in pages_text if p)
    avg_conf = sum(confidences) / len(confidences) if confidences else 0
    warnings = []
    if tables_found:
        warnings.append(f"{tables_found} table(s) extracted")
    if avg_conf < 60:
        warnings.append("Low text confidence - may be scanned PDF")
    return {
        "text": full,
        "pages": len(pages_text),
        "confidence": round(avg_conf, 1),
        "engine": "pdfplumber",
        "warnings": warnings,
    }


def _ocr_pdf_pages(file_path: str) -> dict:
    try:
        from PIL import Image
    except ImportError:
        return _fail("Pillow not installed for image OCR")
    try:
        import pdfplumber
    except ImportError:
        return _fail("pdfplumber not installed")

    pages = []
    confs = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages[:20]):
                img = page.to_image(resolution=200)
                buf = io.BytesIO()
                img.save(buf, format="PNG")
                buf.seek(0)
                pil_img = Image.open(buf)
                ocr_result = _tesseract_ocr_pil(pil_img)
                pages.append(ocr_result["text"])
                confs.append(ocr_result["confidence"])
    except Exception as exc:
        logger.warning("OCR PDF failed: %s", exc)
        return _fail(f"OCR PDF error: {exc}")

    full = "\n\n".join(p for p in pages if p)
    avg = sum(confs) / len(confs) if confs else 0
    return {"text": full, "pages": len(pages), "confidence": round(avg, 1),
            "engine": "pdfplumber+tesseract", "warnings": []}


def _extract_docx(file_path: str) -> dict:
    try:
        import docx
    except ImportError:
        return _fail("python-docx not installed: pip install python-docx")

    doc = docx.Document(file_path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    return {
        "text": text,
        "pages": 1,
        "confidence": 90.0 if text.strip() else 0.0,
        "engine": "python-docx",
        "warnings": [] if text.strip() else ["No text content found in DOCX"],
    }


def _extract_image(file_path: str) -> dict:
    try:
        from PIL import Image
    except ImportError:
        return _fail("Pillow not installed: pip install Pillow")

    img = Image.open(file_path)
    return _tesseract_ocr_pil(img)


def _tesseract_ocr_pil(img) -> dict:
    try:
        import pytesseract
    except ImportError:
        return {"text": "[OCR requires pytesseract - pip install pytesseract + tesseract binary]",
                "pages": 1, "confidence": 0.0, "engine": "pytesseract-missing",
                "warnings": ["pytesseract not installed"]}

    try:
        text = pytesseract.image_to_string(img)
        data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
        confs = [float(c) for c, w in zip(data["conf"], data["text"]) if w.strip() and float(c) > 0]
        avg = sum(confs) / len(confs) if confs else 0
        return {
            "text": text.strip(),
            "pages": 1,
            "confidence": round(avg, 1),
            "engine": "tesseract",
            "warnings": [],
        }
    except Exception as exc:
        return {"text": "", "pages": 1, "confidence": 0.0,
                "engine": "tesseract-error", "warnings": [str(exc)]}


def _fail(msg):
    return {"text": "", "pages": 0, "confidence": 0.0, "engine": "failed", "warnings": [msg]}
